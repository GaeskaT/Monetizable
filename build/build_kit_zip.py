import hashlib
import io
import json
import re
import shutil
import sys
import zipfile
from pathlib import Path

import pymupdf
from docx import Document
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
DELIVERY = ROOT / "delivery"
BUILD = ROOT / "build"
STAGE = BUILD / "zip-stage"
FINAL_SLIDES = BUILD / "final-slides"

KIT = {}
for _index, _arg in enumerate(sys.argv):
    if _arg == "--config" and _index + 1 < len(sys.argv):
        KIT = json.loads(Path(sys.argv[_index + 1]).read_text(encoding="utf-8"))

if KIT:
    PREFIX = KIT["prefix"]
    FOLDER = KIT["zip_stem"]
    TITLE = KIT["title"]
    FIGURE = KIT["figure"]
    PRIMARY_TEXT = KIT["primary_text"]
    FINAL_SLIDES = BUILD / f"final-slides-kit-{KIT['kit']:02d}"
    ZIP_ROOT = ROOT / f"{FOLDER}.zip"
    ZIP_DELIVERY = DELIVERY / f"{FOLDER}.zip"
    REPORT = BUILD / f"kit-{KIT['kit']:02d}-zip-report.json"
    SOURCES = [
        ("00-START-HERE.pdf", DELIVERY / f"{PREFIX}-Start-Here.pdf"),
        ("01-Preaching-Manuscript.pdf", DELIVERY / f"{PREFIX}-Preaching-Manuscript.pdf"),
        ("01-Preaching-Manuscript.docx", DELIVERY / f"{PREFIX}-Preaching-Manuscript.docx"),
        ("02-One-Page-Outline.pdf", DELIVERY / f"{PREFIX}-One-Page-Outline.pdf"),
        ("03-Small-Group-Guide.pdf", DELIVERY / f"{PREFIX}-Small-Group-Guide.pdf"),
        (f"04-Slides-{KIT['deck_stem'].split('-', 3)[3]}.pptx", DELIVERY / f"{KIT['deck_stem']}.pptx"),
        ("06-Preaching-License.pdf", DELIVERY / f"{PREFIX}-Preaching-License.pdf"),
    ]
    PDF_PAGES = {
        "00-START-HERE.pdf": 1,
        "01-Preaching-Manuscript.pdf": KIT["manuscript_pages"],
        "02-One-Page-Outline.pdf": 1,
        "03-Small-Group-Guide.pdf": 1,
        "06-Preaching-License.pdf": 1,
    }
else:
    FOLDER = "Preach-the-Forgotten-Sermon-Kit-1-Mary-of-Bethany"
    TITLE = "What the Room Called Waste"
    FIGURE = "Mary of Bethany"
    PRIMARY_TEXT = "John 12:1–8"
    ZIP_ROOT = ROOT / f"{FOLDER}.zip"
    ZIP_DELIVERY = DELIVERY / f"{FOLDER}.zip"
    REPORT = BUILD / "sermon-kit-zip-report.json"
    SOURCES = [
        ("00-START-HERE.pdf", DELIVERY / "Preach-the-Forgotten-Sermon-Kit-1-Start-Here.pdf"),
        ("01-Preaching-Manuscript.pdf", DELIVERY / "Preach-the-Forgotten-Sermon-Kit-1-Preaching-Manuscript.pdf"),
        ("01-Preaching-Manuscript.docx", DELIVERY / "Preach-the-Forgotten-Sermon-Kit-1-Preaching-Manuscript.docx"),
        ("02-One-Page-Outline.pdf", DELIVERY / "Preach-the-Forgotten-Sermon-Kit-1-One-Page-Outline.pdf"),
        ("03-Small-Group-Guide.pdf", DELIVERY / "Preach-the-Forgotten-Sermon-Kit-1-Small-Group-Guide.pdf"),
        ("04-Slides-What-the-Room-Called-Waste.pptx", DELIVERY / "Sermon-Kit-1-What-the-Room-Called-Waste.pptx"),
        ("06-Preaching-License.pdf", DELIVERY / "Preach-the-Forgotten-Sermon-Kit-1-Preaching-License.pdf"),
    ]
    PDF_PAGES = {
        "00-START-HERE.pdf": 1,
        "01-Preaching-Manuscript.pdf": 19,
        "02-One-Page-Outline.pdf": 1,
        "03-Small-Group-Guide.pdf": 1,
        "06-Preaching-License.pdf": 1,
    }


def fail(msg):
    raise SystemExit(f"ZIP BUILD FAILED: {msg}")


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def build_stage():
    if STAGE.exists():
        shutil.rmtree(STAGE)
    folder = STAGE / FOLDER
    png_dir = folder / "05-Slides-PNG"
    png_dir.mkdir(parents=True)

    for name, source in SOURCES:
        if not source.exists():
            fail(f"missing source: {source}")
        shutil.copy2(source, folder / name)

    pngs = sorted(FINAL_SLIDES.glob("slide-*.png"))
    if len(pngs) != 12:
        fail(f"expected 12 final slide PNGs, got {len(pngs)}")
    for png in pngs:
        shutil.copy2(png, png_dir / png.name)

    return folder


def write_zip(folder):
    entries = []
    for path in sorted(folder.rglob("*")):
        rel = path.relative_to(STAGE).as_posix()
        if path.is_dir():
            entries.append((rel + "/", None))
        else:
            entries.append((rel, path))
    # Keep the folder ahead of its contents and the numbered assets in order.
    entries.sort(key=lambda item: item[0])

    for target in (ZIP_ROOT, ZIP_DELIVERY):
        if target.exists():
            target.unlink()
        with zipfile.ZipFile(target, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as zf:
            for rel, path in entries:
                if path is None:
                    zf.writestr(zipfile.ZipInfo(rel), b"")
                else:
                    zf.write(path, rel)
    return entries


def validate(entries):
    expected = [rel for rel, _ in entries]
    with zipfile.ZipFile(ZIP_ROOT) as zf:
        if zf.testzip() is not None:
            fail("archive failed CRC test")
        actual = zf.namelist()
        if actual != expected:
            fail(f"zip entries do not match staging\nactual={actual}")

        file_names = [n for n in actual if not n.endswith("/")]
        if len(file_names) != 19:
            fail(f"expected 19 files in zip, got {len(file_names)}")

        manifest = {}
        for name, source in SOURCES:
            rel = f"{FOLDER}/{name}"
            data = zf.read(rel)
            manifest[name] = {
                "bytes": len(data),
                "sha256": hashlib.sha256(data).hexdigest(),
            }
            if name.endswith(".pdf"):
                doc = pymupdf.open(stream=data, filetype="pdf")
                pages = doc.page_count
                size = (round(doc[0].rect.width), round(doc[0].rect.height))
                text = doc[0].get_text("text") if pages == 1 else "\n".join(
                    page.get_text("text") for page in doc
                )
                doc.close()
                if pages != PDF_PAGES[name]:
                    fail(f"{name}: expected {PDF_PAGES[name]} pages, got {pages}")
                if size != (612, 792):
                    fail(f"{name}: expected Letter, got {size}")
                if "\ufffd" in text:
                    fail(f"{name}: replacement characters found")
                if "{{" in text and "}}" in text:
                    fail(f"{name}: unresolved template placeholder")
                manifest[name]["pages"] = pages
                manifest[name]["page_size"] = list(size)
                if name == "00-START-HERE.pdf":
                    haystack = text.casefold()
                    for needle in (
                        "From the folder to the pulpit",
                        "Take it to the pulpit",
                        "Guardrails",
                        "The license at a glance",
                        "Tell me how it went",
                    ):
                        if needle.casefold() not in haystack:
                            fail(f"Start Here missing: {needle}")
                if name == "01-Preaching-Manuscript.pdf":
                    haystack = text.casefold()
                    for needle in (
                        TITLE,
                        "Primary text",
                        PRIMARY_TEXT,
                        "Evidence note for the preacher",
                    ):
                        if needle.casefold() not in haystack:
                            fail(f"manuscript missing: {needle}")
            elif name.endswith(".docx"):
                doc = Document(io.BytesIO(data))
                body = "\n".join(p.text for p in doc.paragraphs)
                flat = re.sub(r"\s+", " ", body).casefold()
                if TITLE.casefold() not in flat:
                    fail("DOCX missing title")
                if "evidence note for the preacher" not in flat:
                    fail("DOCX missing evidence note")
                if "{{" in body and "}}" in body:
                    fail("DOCX has unresolved template placeholder")
                manifest[name]["paragraphs"] = len(doc.paragraphs)
            elif name.endswith(".pptx"):
                with zipfile.ZipFile(io.BytesIO(data)) as inner:
                    if inner.testzip() is not None:
                        fail("PPTX inner zip failed CRC test")
                    slides = [
                        n
                        for n in inner.namelist()
                        if n.startswith("ppt/slides/slide") and n.endswith(".xml")
                    ]
                if len(slides) != 12:
                    fail(f"PPTX expected 12 slides, got {len(slides)}")
                manifest[name]["slides"] = len(slides)
                manifest[name]["slide_size_px"] = [1280, 720]

        png_names = sorted(n for n in file_names if n.endswith(".png"))
        if len(png_names) != 12:
            fail(f"expected 12 PNGs, got {len(png_names)}")
        for name in png_names:
            im = Image.open(io.BytesIO(zf.read(name)))
            if im.size != (1280, 720):
                fail(f"{name}: expected 1280x720, got {im.size}")
            manifest[Path(name).name] = {
                "bytes": len(zf.read(name)),
                "size_px": list(im.size),
            }

    if sha256(ZIP_ROOT) != sha256(ZIP_DELIVERY):
        fail("root and delivery zip hashes differ")

    return manifest


def main():
    DELIVERY.mkdir(exist_ok=True)
    folder = build_stage()
    entries = write_zip(folder)
    manifest = validate(entries)
    report = {
        "folder": FOLDER,
        "zip_root": str(ZIP_ROOT),
        "zip_delivery": str(ZIP_DELIVERY),
        "zip_bytes": ZIP_ROOT.stat().st_size,
        "zip_kb": round(ZIP_ROOT.stat().st_size / 1024, 1),
        "zip_sha256": sha256(ZIP_ROOT),
        "file_count": 19,
        "entries": len(entries),
        "files": manifest,
        "checks": {
            "crc": True,
            "entry_order": True,
            "folder_structure": True,
            "file_count": True,
            "pdf_pages_and_size": True,
            "pdf_text": True,
            "docx_readable": True,
            "pptx_12_slides": True,
            "png_1280x720": True,
            "no_placeholders": True,
            "no_replacement_chars": True,
            "root_delivery_match": True,
        },
        "status": "passed",
    }
    REPORT.write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(json.dumps({k: report[k] for k in (
        "zip_root", "zip_delivery", "zip_kb", "zip_sha256", "file_count", "status"
    )}, indent=2))
    print(f"report: {REPORT}")


if __name__ == "__main__":
    main()
