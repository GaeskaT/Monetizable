import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

import pymupdf
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from typography import smarten_paragraph

ROOT = Path(__file__).resolve().parent.parent
BUILD = ROOT / "build"
QA = BUILD / "qa"
DELIVERY = ROOT / "delivery"
SOURCE = ROOT / "SERMON-WHAT-THE-ROOM-CALLED-WASTE.md"
STEM = "Preach-the-Forgotten-Sermon-Kit-1-Preaching-Manuscript"
DOCX = ROOT / f"{STEM}.docx"
PDF = ROOT / f"{STEM}.pdf"
PS1 = BUILD / "_export_manuscript.ps1"

NAVY = RGBColor(0x1B, 0x2A, 0x44)
GOLD = RGBColor(0xB9, 0x8F, 0x4E)
DARK_GOLD = RGBColor(0x8A, 0x6A, 0x2F)
BROWN = RGBColor(0x6A, 0x58, 0x36)
INK = RGBColor(0x26, 0x22, 0x1D)
MUTED = RGBColor(0x6F, 0x66, 0x57)

REQUIRED_TEXT = [
    "What the Room Called Waste",
    "A sermon on Mary of Bethany",
    "PRIMARY TEXT",
    "John 12:1–8",
    "Introduction: The smell in the room",
    "Movement 1: The room full of evidence",
    "Movement 2: The woman who had been listening",
    "Movement 3: The jar that could not be saved",
    "Movement 4: The objection that sounds spiritual",
    "Movement 5: The burial she understood",
    "Application: What do you do with the jar?",
    "1. Name the jar",
    "2. Do one uncalculated act",
    "3. Stop auditing someone else’s devotion",
    "4. Make room for grief",
    "The invitation",
    "Response",
    "Closing",
    "Evidence note for the preacher",
    "May the same be said of us.",
    "Scripture quotations are ESV.",
]

KIT = {}
for _index, _arg in enumerate(sys.argv):
    if _arg == "--config" and _index + 1 < len(sys.argv):
        KIT = json.loads(Path(sys.argv[_index + 1]).read_text(encoding="utf-8"))

if KIT:
    SOURCE = ROOT / KIT["source_manuscript"]
    STEM = f"{KIT['prefix']}-Preaching-Manuscript"
    DOCX = ROOT / f"{STEM}.docx"
    PDF = ROOT / f"{STEM}.pdf"
    REQUIRED_TEXT = KIT["manuscript_required"]
    REPORT = BUILD / f"kit-{KIT['kit']:02d}-manuscript-report.json"
else:
    REPORT = BUILD / "sermon-manuscript-report.json"


def fail(msg):
    print(f"BUILD FAILED: {msg}", file=sys.stderr)
    sys.exit(1)


def normalize(text):
    return re.sub(r"\s+", " ", text).strip()


def compact(text):
    return re.sub(r"\s+", "", text).casefold()


def parse_source():
    raw = SOURCE.read_text(encoding="utf-8")
    if "\ufffd" in raw:
        fail("replacement characters in manuscript source")
    lines = raw.splitlines()
    if not lines or not lines[0].startswith("# "):
        fail("manuscript missing H1")
    title = lines[0][2:].strip()
    first_rule = lines.index("---")
    header_lines = lines[1:first_rule]
    subtitle = ""
    meta = {}
    for line in header_lines:
        stripped = line.strip()
        if not stripped:
            continue
        label_match = re.match(r"^\*\*(.+?):\*\*\s*(.*)$", stripped)
        if label_match:
            meta[label_match.group(1).strip()] = label_match.group(2).strip()
            continue
        subtitle_match = re.match(r"^\*\*(.+?)\*\*$", stripped)
        if subtitle_match:
            subtitle = subtitle_match.group(1).strip()
    body_lines = lines[first_rule + 1 :]
    return title, subtitle, meta, body_lines


def set_character_spacing(run, twips):
    rPr = run._element.get_or_add_rPr()
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:val"), str(twips))
    rPr.insert_element_before(
        spacing,
        "w:w",
        "w:kern",
        "w:position",
        "w:sz",
        "w:szCs",
        "w:highlight",
        "w:u",
        "w:effect",
        "w:bdr",
        "w:shd",
        "w:fitText",
        "w:vertAlign",
        "w:rtl",
        "w:cs",
        "w:em",
        "w:lang",
        "w:eastAsianLayout",
        "w:specVanish",
        "w:oMath",
        "w:rPrChange",
    )


def set_run_font(run, name, size, color, bold=False, italic=False, spacing=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), name)
    if spacing is not None:
        set_character_spacing(run, spacing)


def add_run(paragraph, text, name, size, color, bold=False, italic=False, spacing=None):
    run = paragraph.add_run(text)
    set_run_font(run, name, size, color, bold, italic, spacing)
    return run


def format_paragraph(
    paragraph,
    before=0,
    after=0,
    line=1.0,
    align=None,
    keep_next=False,
    keep_together=False,
):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.keep_with_next = keep_next
    pf.keep_together = keep_together
    pf.widow_control = True
    if align is not None:
        paragraph.alignment = align
    return paragraph


def add_bottom_border(paragraph, color="D9C9A8", size=6, space=2):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.insert_element_before(
        pBdr,
        "w:shd",
        "w:tabs",
        "w:suppressAutoHyphens",
        "w:kinsoku",
        "w:wordWrap",
        "w:overflowPunct",
        "w:topLinePunct",
        "w:autoSpaceDE",
        "w:autoSpaceDN",
        "w:bidi",
        "w:adjustRightInd",
        "w:snapToGrid",
        "w:spacing",
        "w:ind",
        "w:contextualSpacing",
        "w:mirrorIndents",
        "w:suppressOverlap",
        "w:jc",
        "w:textDirection",
        "w:textAlignment",
        "w:textboxTightWrap",
        "w:outlineLvl",
        "w:divId",
        "w:cnfStyle",
        "w:rPr",
        "w:sectPr",
        "w:pPrChange",
    )


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.insert_element_before(
        shd,
        "w:noWrap",
        "w:tcMar",
        "w:textDirection",
        "w:tcFitText",
        "w:vAlign",
        "w:hideMark",
        "w:headers",
        "w:cellIns",
        "w:cellDel",
        "w:cellMerge",
        "w:tcPrChange",
    )


def set_cell_borders(cell, color, size=12, left_only=True):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    sides = ("left",) if left_only else ("top", "left", "bottom", "right")
    for side in sides:
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)
        borders.append(element)
    tcPr.insert_element_before(
        borders,
        "w:shd",
        "w:noWrap",
        "w:tcMar",
        "w:textDirection",
        "w:tcFitText",
        "w:vAlign",
        "w:hideMark",
        "w:headers",
        "w:cellIns",
        "w:cellDel",
        "w:cellMerge",
        "w:tcPrChange",
    )


def set_cell_margins(cell, top=120, bottom=120, left=170, right=170):
    tcPr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for side, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    tcPr.insert_element_before(
        margins,
        "w:textDirection",
        "w:tcFitText",
        "w:vAlign",
        "w:hideMark",
        "w:headers",
        "w:cellIns",
        "w:cellDel",
        "w:cellMerge",
        "w:tcPrChange",
    )


def add_field(paragraph, field_code, name, size, color):
    run = paragraph.add_run()
    set_run_font(run, name, size, color)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = f" {field_code} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instruction)
    run._r.append(separate)
    run._r.append(text)
    run._r.append(end)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Candara"
    normal.font.size = Pt(11.5)
    normal.font.color.rgb = INK
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), "Candara")
    pf = normal.paragraph_format
    pf.line_spacing = 1.32
    pf.space_after = Pt(7)
    pf.space_before = Pt(0)
    pf.widow_control = True

    h2 = doc.styles["Heading 2"]
    h2.font.name = "Georgia"
    h2.font.size = Pt(20)
    h2.font.bold = False
    h2.font.color.rgb = NAVY
    h2.paragraph_format.space_before = Pt(20)
    h2.paragraph_format.space_after = Pt(8)
    h2.paragraph_format.line_spacing = 1.0
    h2.paragraph_format.keep_with_next = True

    h3 = doc.styles["Heading 3"]
    h3.font.name = "Georgia"
    h3.font.size = Pt(12.5)
    h3.font.bold = True
    h3.font.color.rgb = NAVY
    h3.paragraph_format.space_before = Pt(12)
    h3.paragraph_format.space_after = Pt(3)
    h3.paragraph_format.line_spacing = 1.0
    h3.paragraph_format.keep_with_next = True

    bullet = doc.styles["List Bullet"]
    bullet.font.name = "Candara"
    bullet.font.size = Pt(10.5)
    bullet.font.color.rgb = INK
    bullet.paragraph_format.space_after = Pt(3)
    bullet.paragraph_format.line_spacing = 1.25


def build_title_page(doc, title, subtitle, meta):
    content_width = Inches(6.45)

    masthead = doc.add_paragraph()
    format_paragraph(masthead, after=0, line=1.0)
    masthead.paragraph_format.tab_stops.add_tab_stop(content_width, WD_TAB_ALIGNMENT.RIGHT)
    add_run(masthead, "SILENCED SAINTS", "Candara", 8, DARK_GOLD, bold=True, spacing=34)
    add_run(masthead, "\t" + KIT.get("masthead_right", "PREACH THE FORGOTTEN · SERMON KIT 1"), "Candara", 7.5, NAVY, spacing=14)
    add_bottom_border(masthead, color="C8A15E", size=8, space=4)

    for _ in range(2):
        spacer = doc.add_paragraph()
        format_paragraph(spacer, after=0, line=1.0)
        add_run(spacer, "", "Candara", 20, INK)

    kicker = doc.add_paragraph()
    format_paragraph(kicker, before=30, after=3, line=1.0)
    add_run(kicker, "A SERMON FOR THE PULPIT", "Candara", 8.5, DARK_GOLD, bold=True, spacing=40)

    title_p = doc.add_paragraph()
    format_paragraph(title_p, before=0, after=6, line=1.02)
    title_lines = title.split(" ", 3)
    display_lines = KIT.get("display_lines", ["What the Room", "Called Waste"])
    for index, line in enumerate(display_lines):
        if index:
            title_p.add_run().add_break()
        add_run(title_p, line, "Georgia", 40, NAVY)

    subtitle_p = doc.add_paragraph()
    format_paragraph(subtitle_p, before=2, after=0, line=1.2)
    add_run(subtitle_p, subtitle, "Georgia", 15, BROWN, italic=True)

    rule = doc.add_paragraph()
    format_paragraph(rule, before=10, after=16, line=1.0)
    add_bottom_border(rule, color="C8A15E", size=16, space=1)

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.45)
    shade_cell(cell, "FAF6EC")
    set_cell_borders(cell, "C8A15E", size=24, left_only=True)
    set_cell_margins(cell, top=150, bottom=150, left=210, right=210)
    label_p = cell.paragraphs[0]
    format_paragraph(label_p, after=4, line=1.0)
    add_run(label_p, "THE BIG IDEA", "Candara", 7.5, DARK_GOLD, bold=True, spacing=30)
    idea_p = cell.add_paragraph()
    format_paragraph(idea_p, before=0, after=0, line=1.22)
    add_run(
        idea_p,
        smarten_paragraph(meta.get("Big idea", "")),
        "Georgia",
        12.5,
        NAVY,
        italic=True,
    )

    meta_p = doc.add_paragraph()
    format_paragraph(meta_p, before=18, after=0, line=1.0)
    meta_p.paragraph_format.tab_stops.add_tab_stop(Inches(1.25))
    rows = [
        ("PRIMARY TEXT", meta.get("Primary text", "")),
        ("ALSO READ", meta.get("Also read", "")),
        ("TIME", meta.get("Time", "")),
    ]
    for index, (label, value) in enumerate(rows):
        if index:
            meta_p.add_run().add_break()
        add_run(meta_p, label, "Candara", 7.2, DARK_GOLD, bold=True, spacing=20)
        add_run(meta_p, "\t" + smarten_paragraph(value), "Candara", 9.2, NAVY)

    note_p = doc.add_paragraph()
    format_paragraph(note_p, before=8, after=0, line=1.2)
    note_label = meta.get("Note", "")
    add_run(note_p, smarten_paragraph(note_label), "Candara", 8.4, MUTED, italic=True)

    bottom_spacer = doc.add_paragraph()
    format_paragraph(bottom_spacer, before=72, after=0, line=1.0)
    add_run(bottom_spacer, "", "Candara", 8, INK)

    seal_rule = doc.add_paragraph()
    format_paragraph(seal_rule, before=0, after=6, line=1.0)
    add_bottom_border(seal_rule, color="D9C9A8", size=6, space=1)
    add_run(seal_rule, "A COMPLETE MANUSCRIPT FOR THE PULPIT", "Candara", 7.6, DARK_GOLD, bold=True, spacing=30)
    edition_p = doc.add_paragraph()
    format_paragraph(edition_p, before=2, after=0, line=1.0)
    add_run(edition_p, "FIRST EDITION · OCTOBER 2026", "Candara", 7.4, MUTED, spacing=20)

    doc.add_page_break()


def build_header_footer(section):
    section.different_first_page_header_footer = True

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    format_paragraph(hp, after=0, line=1.0)
    hp.paragraph_format.tab_stops.add_tab_stop(Inches(6.45), WD_TAB_ALIGNMENT.RIGHT)
    add_run(hp, KIT.get("header_title", "What the Room Called Waste"), "Candara", 7.8, BROWN, spacing=10)
    add_run(hp, "\t" + KIT.get("header_right", "Mary of Bethany · Sermon Kit 1"), "Candara", 7.4, DARK_GOLD, spacing=8)
    add_bottom_border(hp, color="D9C9A8", size=4, space=2)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    format_paragraph(fp, after=0, line=1.0)
    fp.paragraph_format.tab_stops.add_tab_stop(Inches(6.45), WD_TAB_ALIGNMENT.RIGHT)
    add_run(fp, KIT.get("footer_left", "PREACH THE FORGOTTEN · SERMON KIT 1"), "Candara", 6.6, DARK_GOLD, spacing=20)
    add_run(fp, "\tPAGE ", "Candara", 6.6, NAVY, spacing=12)
    add_field(fp, "PAGE", "Candara", 6.8, NAVY)


def heading(doc, text, level):
    p = doc.add_paragraph(style="Heading 2" if level == 2 else "Heading 3")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    if level == 2:
        add_run(p, smarten_paragraph(text), "Georgia", 20, NAVY)
        add_bottom_border(p, color="D9C9A8", size=6, space=3)
    else:
        add_run(p, smarten_paragraph(text), "Georgia", 12.5, NAVY, bold=True)
    return p


def body(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p, before=0, after=7, line=1.32)
    add_run(p, smarten_paragraph(text), "Candara", 11.5, INK)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    format_paragraph(p, before=0, after=3, line=1.25)
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    add_run(p, smarten_paragraph(text), "Candara", 10.5, INK)
    return p


def build_body(doc, body_lines):
    in_evidence = False
    for raw_line in body_lines:
        line = raw_line.rstrip()
        if not line or line == "---":
            continue
        if line.startswith("## "):
            text = line[3:].strip()
            heading(doc, text, 2)
            in_evidence = text.startswith("Evidence note")
            continue
        if line.startswith("### "):
            heading(doc, line[4:].strip(), 3)
            continue
        if line.startswith("- "):
            bullet(doc, line[2:].strip())
            continue
        body(doc, line.strip())


def set_core_properties(doc):
    props = doc.core_properties
    props.title = KIT.get("prop_title", "What the Room Called Waste - Preaching Manuscript")
    props.subject = KIT.get("prop_subject", "Sermon manuscript for Mary of Bethany, Sermon Kit 1")
    props.author = "Silenced Saints"
    props.last_modified_by = "Silenced Saints"
    props.keywords = KIT.get(
        "keywords",
        "Mary of Bethany, What the Room Called Waste, Preach the Forgotten, "
        "Sermon Kit 1, preaching manuscript, Silenced Saints",
    )
    props.comments = KIT.get(
        "comments",
        "Editable preaching manuscript. ESV quotations. One-page outline and "
        "small-group guide sold separately in Sermon Kit 1.",
    )
    props.category = "Sermon Kit"


def build_docx():
    title, subtitle, meta, body_lines = parse_source()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.05)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)

    configure_styles(doc)
    build_title_page(doc, title, subtitle, meta)
    build_header_footer(section)
    build_body(doc, body_lines)
    set_core_properties(doc)
    doc.save(DOCX)

    check = Document(DOCX)
    h2 = sum(1 for p in check.paragraphs if p.style.name == "Heading 2")
    h3 = sum(1 for p in check.paragraphs if p.style.name == "Heading 3")
    words = sum(len(p.text.split()) for p in check.paragraphs)
    if h2 < 11:
        fail(f"DOCX only has {h2} H2 headings")
    if h3 < 4:
        fail(f"DOCX only has {h3} H3 headings")
    if words < 6800:
        fail(f"DOCX word count too low: {words}")
    if "\ufffd" in "\n".join(p.text for p in check.paragraphs):
        fail("DOCX contains replacement characters")
    return {"h2": h2, "h3": h3, "words": words, "docx_kb": round(DOCX.stat().st_size / 1024, 1)}


def write_ps1():
    PS1.write_text(
        r'''param([string]$Docx,[string]$Pdf)
$ErrorActionPreference = "Stop"
$word = New-Object -ComObject Word.Application
$word.Visible = $false
$word.DisplayAlerts = 0
$doc = $word.Documents.Open($Docx, $false, $false)
try {
  $doc.Fields.Update() | Out-Null
  $pages = $doc.ComputeStatistics(2)
  if (Test-Path $Pdf) { Remove-Item $Pdf -Force }
  $doc.SaveAs([ref]$Pdf, [ref]17)
  Write-Output ("PAGES=" + $pages)
} finally {
  $doc.Close($false)
  $word.Quit()
}
''',
        encoding="utf-8",
    )


def export_pdf():
    write_ps1()
    proc = subprocess.run(
        [
            "powershell",
            "-NoProfile",
            "-ExecutionPolicy",
            "Bypass",
            "-File",
            str(PS1),
            "-Docx",
            str(DOCX),
            "-Pdf",
            str(PDF),
        ],
        capture_output=True,
        text=True,
        timeout=300,
    )
    output = (proc.stdout or "") + (proc.stderr or "")
    if proc.returncode != 0 or not PDF.exists():
        fail(f"Word export failed\n{output[-3000:]}")
    page_match = re.search(r"PAGES=(\d+)", output)
    return int(page_match.group(1)) if page_match else None


def verify_pdf(word_pages=None):
    doc = pymupdf.open(PDF)
    if doc.page_count < 12:
        doc.close()
        fail(f"PDF only has {doc.page_count} pages")
    if word_pages and doc.page_count != word_pages:
        doc.close()
        fail(f"PDF has {doc.page_count} pages but Word reported {word_pages}")
    for index in range(doc.page_count):
        page = doc[index]
        width, height = round(page.rect.width), round(page.rect.height)
        if (width, height) != (612, 792):
            doc.close()
            fail(f"page {index + 1} is {width}x{height}, not Letter")

    pages_text = [page.get_text("text") for page in doc]
    full = normalize("\n".join(pages_text))
    squeezed = compact(full)
    missing = [text for text in REQUIRED_TEXT if compact(text) not in squeezed]
    if missing:
        doc.close()
        fail(f"PDF missing text: {missing}")
    if "\ufffd" in full:
        doc.close()
        fail("PDF contains replacement characters")

    fonts = sorted({font[3] for page in doc for font in page.get_fonts(full=True)})
    lowered = " ".join(fonts).lower()
    for family in ("georgia", "candara"):
        if family not in lowered:
            doc.close()
            fail(f"{family} not embedded; fonts={fonts}")

    bounds = []
    for page in doc:
        blocks = page.get_text("blocks")
        if not blocks:
            continue
        bounds.append(
            (
                min(b[0] for b in blocks),
                max(b[2] for b in blocks),
                max(b[3] for b in blocks),
            )
        )
    min_x = min(b[0] for b in bounds)
    max_x = max(b[1] for b in bounds)
    max_y = max(b[2] for b in bounds)
    if min_x < 48 or max_x > 566:
        doc.close()
        fail(f"text outside horizontal safe area: {min_x:.1f}-{max_x:.1f}")
    if max_y > 772:
        doc.close()
        fail(f"text reaches y={max_y:.1f}")
    if any("" == page.get_text("text").strip() for page in doc):
        doc.close()
        fail("PDF contains a blank page")

    doc.close()
    return {
        "pages": len(pages_text),
        "words": len(full.split()),
        "fonts": fonts,
        "text_bounds_x": [round(min_x, 1), round(max_x, 1)],
        "max_text_y": round(max_y, 1),
        "pdf_kb": round(PDF.stat().st_size / 1024, 1),
    }


def render_previews(page_count):
    QA.mkdir(exist_ok=True)
    doc = pymupdf.open(PDF)
    targets = {"manuscript-title": 0, "manuscript-body": 3, "manuscript-close": page_count - 1}
    out = {}
    for name, index in targets.items():
        pix = doc[index].get_pixmap(dpi=110)
        path = QA / f"{name}.png"
        pix.save(path)
        out[name] = str(path)
    doc.close()
    return out


def main():
    DELIVERY.mkdir(exist_ok=True)
    info = build_docx()
    word_pages = export_pdf()
    pdf_info = verify_pdf(word_pages)
    previews = render_previews(pdf_info["pages"])
    pdf_delivery = DELIVERY / PDF.name
    docx_delivery = DELIVERY / DOCX.name
    shutil.copy2(PDF, pdf_delivery)
    shutil.copy2(DOCX, docx_delivery)
    report = {
        "docx": str(DOCX),
        "pdf": str(PDF),
        "delivery_pdf": str(pdf_delivery),
        "delivery_docx": str(docx_delivery),
        **info,
        **pdf_info,
        "word_pages": word_pages,
        "previews": previews,
        "status": "passed",
    }
    out = REPORT
    out.write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
