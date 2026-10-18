import json
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
BUILD = ROOT / "build"
FONTS = Path("C:/Windows/Fonts")

NAVY = "101B2E"
NAVY_DEEP = "0C1727"
NAVY_BOTTOM = "0A1423"
CREAM = "F4ECDC"
CREAM_SOFT = "D8C5A1"
GOLD = "C8A15E"
MUTED = "B9AB90"
SLATE = "8F9AAB"
FOOT = "66748B"
INK = "16233A"
BROWN = "6A5836"
WHITE_LINE = "D9C9A8"

KIT = json.loads(Path(sys.argv[sys.argv.index("--config") + 1]).read_text(encoding="utf-8"))
DECK = KIT["deck"]
SLIDES = DECK["slides"]
OUT_PPTX = ROOT / f"{KIT['deck_stem']}.pptx"
REPORT = BUILD / f"kit-{KIT['kit']:02d}-deck-report.json"


def rgb(value):
    return RGBColor.from_string(value)


def textbox(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for index, spec in enumerate(paras):
        para = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        para.alignment = spec.get("align", PP_ALIGN.LEFT)
        para.line_spacing = spec.get("line", 1.0)
        para.space_before = Pt(0)
        para.space_after = Pt(0)
        for run_spec in spec["runs"]:
            text, name, size, color = run_spec[:4]
            bold = run_spec[4] if len(run_spec) > 4 else False
            run = para.add_run()
            run.text = text
            run.font.name = name
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = rgb(color)
    return box


def rect(slide, x, y, w, h, fill, alpha=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.fill.background()
    shape.shadow.inherit = False
    if alpha is not None:
        solid = shape.fill._xPr.find(
            "{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill"
        )
        color = solid[0]
        node = etree.SubElement(
            color, "{http://schemas.openxmlformats.org/drawingml/2006/main}alpha"
        )
        node.set("val", str(int(alpha * 100000)))
    return shape


def background(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def footer(slide, number, right=None):
    textbox(
        slide,
        0.79,
        7.04,
        7.29,
        0.21,
        [{"runs": [(KIT["footer_left"], "Candara", 8.5, FOOT)]}],
    )
    textbox(
        slide,
        11.46,
        7.04,
        1.08,
        0.21,
        [
            {
                "align": PP_ALIGN.RIGHT,
                "runs": [(right or f"{number:02d} / 12", "Candara", 9, GOLD, True)],
            }
        ],
    )


def header(slide, label, ref):
    textbox(
        slide,
        0.79,
        0.54,
        7.4,
        0.25,
        [{"runs": [(label, "Candara", 10.5, GOLD, True)]}],
    )
    if ref:
        textbox(
            slide,
            8.0,
            0.54,
            4.54,
            0.23,
            [{"align": PP_ALIGN.RIGHT, "runs": [(ref, "Candara", 9, SLATE)]}],
        )


def headline(slide, text, size=42, y=1.0, h=1.0):
    textbox(
        slide,
        0.79,
        y,
        11.75,
        h,
        [{"runs": [(text, "Georgia", size, CREAM)]}],
    )


def new_slide(prs, color=NAVY):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide, color)
    return slide


def slide_title(prs, spec):
    slide = new_slide(prs, NAVY_DEEP)
    slide.shapes.add_picture(
        str(ROOT / KIT["cover_art"]), Inches(8.33), 0, Inches(5.0), Inches(7.5)
    )
    textbox(
        slide,
        0.79,
        0.54,
        6.77,
        0.29,
        [{"runs": [("SILENCED SAINTS", "Candara", 12, GOLD, True)]}],
    )
    textbox(
        slide,
        0.79,
        0.90,
        6.77,
        0.23,
        [{"runs": [(spec["kicker"], "Candara", 9.5, CREAM_SOFT)]}],
    )
    lines = spec["title_lines"]
    textbox(
        slide,
        0.77,
        1.92,
        7.4,
        2.81,
        [
            {"line": 0.94, "runs": [(line, "Georgia", spec.get("title_size", 72), CREAM)]}
            for line in lines
        ],
    )
    textbox(
        slide,
        0.79,
        4.94,
        6.77,
        0.56,
        [{"runs": [(spec["figure"], "Georgia", 30, GOLD)]}],
    )
    textbox(
        slide,
        0.79,
        5.56,
        6.77,
        0.29,
        [{"runs": [(spec["text"], "Candara", 14, CREAM, True)]}],
    )
    textbox(
        slide,
        0.79,
        6.42,
        6.77,
        0.25,
        [{"runs": [(spec["run"], "Candara", 9, MUTED)]}],
    )
    textbox(
        slide,
        0.79,
        6.73,
        6.77,
        0.21,
        [{"runs": [("FIRST EDITION · OCTOBER 2026", "Candara", 8.5, SLATE)]}],
    )


def slide_idea(prs, spec, number):
    slide = new_slide(prs)
    header(slide, spec["label"], spec["ref"])
    textbox(
        slide,
        0.79,
        1.83,
        11.75,
        2.73,
        [
            {"line": 1.05, "runs": [(line, "Georgia", spec.get("size", 52), CREAM)]}
            for line in spec["lines"]
        ],
    )
    textbox(
        slide,
        0.79,
        4.81,
        8.54,
        0.38,
        [{"runs": [(spec["note"], "Candara", 20, MUTED)]}],
    )
    footer(slide, number)


def slide_columns(prs, spec, number, widths):
    slide = new_slide(prs)
    header(slide, spec["label"], spec["ref"])
    headline(slide, spec["headline"], spec.get("headline_size", 42))
    xs = [0.79 + i * widths["step"] for i in range(len(spec["columns"]))]
    for x, col in zip(xs, spec["columns"]):
        textbox(
            slide,
            x,
            2.15,
            widths["col"],
            0.21,
            [{"runs": [(col["num"], "Candara", 11, GOLD, True)]}],
        )
        textbox(
            slide,
            x,
            2.46,
            widths["col"],
            0.67,
            [{"runs": [(col["name"], "Georgia", widths["name_size"], CREAM)]}],
        )
        textbox(
            slide,
            x,
            3.23,
            widths["col"],
            0.83,
            [
                {"line": 1.1, "runs": [(line, "Candara", widths["body_size"], MUTED)]}
                for line in col["desc"]
            ],
        )
    rect(slide, 0.79, 4.69, 11.75, 1.33, CREAM)
    textbox(
        slide,
        1.12,
        4.93,
        11.08,
        0.80,
        [
            {"runs": [(line, "Georgia", 24, INK)]}
            for line in spec["panel"]
        ],
        anchor=MSO_ANCHOR.MIDDLE,
    )
    textbox(
        slide,
        0.79,
        6.27,
        8.54,
        0.29,
        [{"runs": [(spec["kicker"], "Candara", 16, GOLD)]}],
    )
    footer(slide, number)


def slide_three(prs, spec, number):
    slide = new_slide(prs)
    header(slide, spec["label"], spec["ref"])
    headline(slide, spec["headline"], spec.get("headline_size", 46))
    xs = [0.79, 4.79, 8.79]
    for x, col in zip(xs, spec["columns"]):
        textbox(
            slide,
            x,
            1.98,
            3.75,
            0.21,
            [{"runs": [(col["num"], "Candara", 11, GOLD, True)]}],
        )
        textbox(
            slide,
            x,
            2.25,
            3.75,
            0.23,
            [{"runs": [(col["ref"], "Candara", 11, CREAM_SOFT, True)]}],
        )
        textbox(
            slide,
            x,
            2.56,
            3.75,
            1.35,
            [{"runs": [(line, "Georgia", 34, CREAM if i == 0 else GOLD)]} for i, line in enumerate(col["lines"])],
        )
        textbox(
            slide,
            x,
            4.08,
            3.75,
            0.62,
            [{"line": 1.1, "runs": [(col["note"], "Candara", 15, MUTED)]}],
        )
    textbox(
        slide,
        0.79,
        5.31,
        11.75,
        1.04,
        [
            {"line": 1.1, "runs": [(line, "Georgia", 28, CREAM)]}
            for line in spec["panel"]
        ],
    )
    footer(slide, number)


def slide_stat(prs, spec, number):
    slide = new_slide(prs)
    header(slide, spec["label"], spec["ref"])
    headline(slide, spec["headline"], spec.get("headline_size", 42))
    textbox(
        slide,
        0.79,
        1.88,
        5.0,
        3.12,
        [{"line": 0.9, "runs": [(spec["big"], "Georgia", spec.get("big_size", 156), GOLD)]}],
    )
    textbox(
        slide,
        0.85,
        5.0,
        4.79,
        0.56,
        [{"runs": [(spec["big_label"], "Georgia", 34, CREAM)]}],
    )
    textbox(
        slide,
        0.85,
        5.56,
        4.79,
        0.27,
        [{"runs": [(spec["big_sub"], "Candara", 14, MUTED)]}],
    )
    textbox(
        slide,
        6.25,
        1.98,
        6.29,
        1.88,
        [{"line": 1.25, "runs": [(spec["body"], "Candara", 18, CREAM)]}],
    )
    textbox(
        slide,
        6.25,
        4.06,
        6.29,
        1.98,
        [{"line": 1.12, "runs": [(spec["pullquote"], "Georgia", 30, GOLD)]}],
    )
    textbox(
        slide,
        0.79,
        6.35,
        10.42,
        0.38,
        [{"runs": [(spec["footline"], "Georgia", 20, MUTED)]}],
    )
    footer(slide, number)


def slide_objection(prs, spec, number):
    slide = new_slide(prs)
    header(slide, spec["label"], spec["ref"])
    headline(slide, spec["headline"], spec.get("headline_size", 42))
    textbox(
        slide,
        0.79,
        1.88,
        5.62,
        1.88,
        [{"line": 1.05, "runs": [(spec["quote"], "Georgia", spec.get("quote_size", 24), CREAM)]}],
    )
    textbox(
        slide,
        0.79,
        3.77,
        5.62,
        0.23,
        [{"runs": [(spec["quote_ref"], "Candara", 11, GOLD, True)]}],
    )
    textbox(
        slide,
        6.88,
        1.98,
        5.67,
        1.15,
        [{"line": 1.05, "runs": [(spec["aside1"], "Georgia", 24, GOLD)]}],
    )
    textbox(
        slide,
        6.88,
        3.25,
        5.67,
        0.83,
        [{"line": 1.1, "runs": [(spec["aside2"], "Candara", 16, MUTED)]}],
    )
    for x, box in ((0.79, spec["box1"]), (6.88, spec["box2"])):
        rect(slide, x, 4.06, 5.67, 2.08, CREAM)
        textbox(
            slide,
            x + 0.29,
            4.31,
            5.08,
            0.23,
            [{"runs": [(box["label"], "Candara", 11, BROWN, True)]}],
        )
        textbox(
            slide,
            x + 0.29,
            4.56,
            3.12,
            1.27,
            [{"runs": [(box["big"], "Georgia", 66, INK)]}],
        )
        textbox(
            slide,
            x + 0.29,
            5.85,
            5.08,
            0.25,
            [{"runs": [(box["sub"], "Candara", 14, BROWN)]}],
        )
    textbox(
        slide,
        0.79,
        6.31,
        11.75,
        0.58,
        [{"line": 1.05, "runs": [(spec["footline"], "Candara", 16, CREAM_SOFT)]}],
    )
    footer(slide, number)


def slide_verdict(prs, spec, number):
    slide = new_slide(prs)
    header(slide, spec["label"], spec["ref"])
    textbox(
        slide,
        0.79,
        1.98,
        11.75,
        1.25,
        [{"line": 1.05, "runs": [(spec["line1"], "Georgia", 60, "7F8B9F")]}],
    )
    textbox(
        slide,
        0.79,
        3.33,
        11.75,
        1.35,
        [{"line": 1.05, "runs": [(spec["line2"], "Georgia", 60, GOLD)]}],
    )
    textbox(
        slide,
        0.79,
        5.1,
        9.9,
        0.42,
        [{"runs": [(spec["footline"], "Candara", 20, MUTED)]}],
    )
    footer(slide, number)


def slide_banner(prs, spec, number, kind):
    slide = new_slide(prs)
    header(slide, spec["label"], spec["ref"])
    rect(slide, 0.79, 1.25, 11.75, 2.6, CREAM)
    textbox(
        slide,
        1.12,
        1.67,
        11.08,
        1.04,
        [
            {
                "align": PP_ALIGN.CENTER,
                "runs": [(spec["banner_main"], "Georgia", spec.get("banner_size", 44), INK)],
            }
        ],
    )
    textbox(
        slide,
        1.12,
        2.83,
        11.08,
        0.33,
        [
            {
                "align": PP_ALIGN.CENTER,
                "runs": [(spec["banner_sub"], "Candara", 18, BROWN)],
            }
        ],
    )
    if kind == "scripture":
        rect(slide, 11.85, 2.72, 0.08, 1.12, GOLD)
        rect(slide, 11.35, 3.17, 1.08, 0.08, GOLD)
    textbox(
        slide,
        0.79,
        4.48,
        11.75,
        1.56,
        [
            {"line": 1.08, "runs": [(line, "Georgia", 30, color)]}
            for line, color in zip(spec["lines"], (CREAM, GOLD))
        ],
    )
    textbox(
        slide,
        0.79,
        6.15,
        9.9,
        0.31,
        [{"runs": [(spec["footline"], "Candara", 17, MUTED)]}],
    )
    footer(slide, number)


def slide_response(prs, spec, number):
    slide = new_slide(prs)
    header(slide, spec["label"], spec["ref"])
    headline(slide, spec["headline"], spec.get("headline_size", 42))
    positions = [(0.79, 1.98), (6.96, 1.98), (0.79, 4.17), (6.96, 4.17)]
    for (x, y), item in zip(positions, spec["items"]):
        textbox(
            slide,
            x,
            y,
            5.58,
            0.21,
            [{"runs": [(item["num"], "Candara", 11, GOLD, True)]}],
        )
        textbox(
            slide,
            x,
            y + 0.27,
            5.58,
            1.0,
            [{"line": 1.0, "runs": [(line, "Georgia", 30, CREAM)]} for line in item["title"]],
        )
        textbox(
            slide,
            x,
            y + 1.31,
            5.58,
            0.72,
            [{"line": 1.1, "runs": [(item["desc"], "Candara", 15, MUTED)]}],
        )
    textbox(
        slide,
        0.79,
        6.46,
        9.38,
        0.38,
        [{"runs": [(spec["footline"], "Georgia", 20, GOLD)]}],
    )
    footer(slide, number)


def slide_invitation(prs, spec, number):
    slide = new_slide(prs)
    header(slide, spec["label"], spec["ref"])
    textbox(
        slide,
        0.79,
        1.67,
        11.75,
        2.29,
        [{"line": 1.0, "runs": [(spec["big"], "Georgia", spec.get("big_size", 60), CREAM)]}],
    )
    textbox(
        slide,
        1.88,
        4.27,
        9.58,
        1.04,
        [
            {
                "align": PP_ALIGN.CENTER,
                "line": 1.1,
                "runs": [(spec["line1"], "Georgia", 24, GOLD)],
            }
        ],
    )
    textbox(
        slide,
        1.46,
        5.62,
        10.42,
        1.04,
        [
            {
                "align": PP_ALIGN.CENTER,
                "line": 1.1,
                "runs": [(line, "Candara", 26, CREAM)],
            }
            for line in spec["lines"]
        ],
    )
    footer(slide, number)


def slide_closing(prs, spec, number):
    slide = new_slide(prs, NAVY_BOTTOM)
    picture = slide.shapes.add_picture(
        str(ROOT / KIT["cover_art"]), 0, 0, Inches(13.3333), Inches(7.5)
    )
    picture.crop_left = 0.31131
    picture.crop_right = 0.31131
    rect(slide, 0, 0, 13.3333, 7.5, NAVY_BOTTOM, alpha=0.80)
    textbox(
        slide,
        1.46,
        0.54,
        10.42,
        0.25,
        [
            {
                "align": PP_ALIGN.CENTER,
                "runs": [(spec["kicker"], "Candara", 9, GOLD)],
            }
        ],
    )
    textbox(
        slide,
        1.46,
        1.83,
        10.42,
        2.6,
        [
            {"align": PP_ALIGN.CENTER, "line": 1.05, "runs": [(line, "Georgia", 54, CREAM)]}
            for line in spec["lines"]
        ],
    )
    textbox(
        slide,
        1.46,
        4.65,
        10.42,
        0.52,
        [
            {
                "align": PP_ALIGN.CENTER,
                "runs": [(spec["line2"], "Georgia", 22, GOLD)],
            }
        ],
    )
    textbox(
        slide,
        1.46,
        5.42,
        10.42,
        0.83,
        [
            {
                "align": PP_ALIGN.CENTER,
                "line": 1.2,
                "runs": [(spec["body"], "Candara", 20, CREAM)],
            }
        ],
    )
    textbox(
        slide,
        1.46,
        6.73,
        10.42,
        0.23,
        [
            {
                "align": PP_ALIGN.CENTER,
                "runs": [("FIRST EDITION · OCTOBER 2026", "Candara", 8.5, SLATE)],
            }
        ],
    )
    textbox(
        slide,
        11.46,
        7.04,
        1.08,
        0.21,
        [{"align": PP_ALIGN.RIGHT, "runs": [("12 / 12", "Candara", 9, GOLD, True)]}],
    )


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.3333)
    prs.slide_height = Inches(7.5)

    for index, spec in enumerate(SLIDES, start=1):
        layout = spec["layout"]
        if layout == "title":
            slide_title(prs, spec)
        elif layout == "idea":
            slide_idea(prs, spec, index)
        elif layout == "columns":
            slide_columns(
                prs,
                spec,
                index,
                {
                    "step": 3.02,
                    "col": 2.71,
                    "name_size": spec.get("name_size", 34),
                    "body_size": spec.get("body_size", 15),
                },
            )
        elif layout == "three":
            slide_three(prs, spec, index)
        elif layout == "stat":
            slide_stat(prs, spec, index)
        elif layout == "objection":
            slide_objection(prs, spec, index)
        elif layout == "verdict":
            slide_verdict(prs, spec, index)
        elif layout == "scripture":
            slide_banner(prs, spec, index, "scripture")
        elif layout == "gospel":
            slide_banner(prs, spec, index, "gospel")
        elif layout == "response":
            slide_response(prs, spec, index)
        elif layout == "invitation":
            slide_invitation(prs, spec, index)
        elif layout == "closing":
            slide_closing(prs, spec, index)
        else:
            raise SystemExit(f"unknown layout: {layout}")

    prs.core_properties.title = KIT["listing_title"]
    prs.core_properties.author = "Silenced Saints"
    prs.core_properties.subject = KIT["prop_subject"]
    prs.core_properties.keywords = KIT["keywords"]
    prs.core_properties.comments = KIT["comments"]
    prs.core_properties.category = "Sermon slides"
    prs.save(OUT_PPTX)

    report = {
        "file": str(OUT_PPTX),
        "slides": len(prs.slides),
        "size_px": [1280, 720],
        "slide_size_in": [13.3333, 7.5],
        "fonts": ["Georgia", "Candara"],
        "source_art": KIT["cover_art"],
        "status": "passed",
    }
    REPORT.write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    build()
